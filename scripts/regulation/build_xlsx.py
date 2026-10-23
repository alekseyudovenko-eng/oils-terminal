"""Собирает Excel-таблицу по регуляторной политике из data/regulation.json.

Запуск: python scripts/regulation/build_xlsx.py [--lang=en] [путь_к_xlsx]
По умолчанию пишет public/lab/regulation.xlsx (русская версия) или public/lab/regulation-en.xlsx (английская);
обе закрыты паролем вместе с /lab.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
DATA = json.loads((ROOT / 'data' / 'regulation.json').read_text(encoding='utf-8'))

L = {
    'ru': {
        'bloc': {'EU': 'ЕС', 'EAEU': 'ЕАЭС', 'other': 'Вне ЕС и ЕАЭС'},
        'use': {'high': 'Высокая', 'medium': 'Средняя', 'low': 'Низкая', 'none': 'Нет', 'nodata': 'Нет данных'},
        'type': {
            'disclosure': 'Раскрытие в составе / маркировка', 'claims': 'Правила заявок',
            'ban': 'Запрет / ограничение состава', 'dairy': 'Молочка и заменители молочного жира',
            'import': 'Импортный контроль', 'traceability': 'Цифровая маркировка и прослеживаемость',
            'tax': 'Налоги', 'standard': 'Стандарты качества', 'sustainability': 'Устойчивость (EUDR)',
        },
        'status': {'in_force': 'Действует', 'upcoming': 'Принято, вступает позже', 'proposal': 'Предложение / проект'},
        'scope': {'EU': 'ЕС (общие нормы)', 'EAEU': 'ЕАЭС (общие нормы)'},
        'sheets': ('Сводка по странам', 'Меры', 'Справочник'),
        'mcols': [('№', 5), ('Код охвата', 9), ('Охват', 20), ('Тип меры', 24), ('Статус', 17), ('Дата вступления', 13),
                  ('Правовой акт / источник нормы', 34), ('Мера', 34), ('Описание', 60),
                  ('Measure (EN)', 34), ('Description (EN)', 60), ('Ссылка', 40)],
        'ccols': [('№', 5), ('Страна', 16), ('Блок', 14), ('Формулировка на упаковке', 26),
                  ('Использование заявки «без пальмового масла»', 16), ('Норм всего (вкл. блок)', 11), ('Действуют', 11),
                  ('Вступают позже', 11), ('Предложения', 11), ('Регулирование кратко', 52), ('Рыночная практика заявки', 60),
                  ('Код', 7), ('Код блока', 8)],
        'total': 'Итого норм в базе',
        'note': 'Столбцы F–I считаются формулами по листу «Меры»: нормы страны + общие нормы её блока (ЕС или ЕАЭС). '
                'Столбец E — экспертная оценка по публикациям, выбирается из списка.',
        'guide': [
            ('Регуляторная политика: маркировки и пальмовое масло', ''),
            ('Подготовлено', 'International Marketing Agency (IMA)'),
            ('Данные на', None), ('Период обзора', None), ('', ''), ('Статусы', ''), 'STATUS', ('', ''), ('Типы мер', ''), 'TYPE', ('', ''),
            ('Шкала использования заявки', 'Экспертная оценка по найденным публикациям: Высокая / Средняя / Низкая / Нет / Нет данных.'),
            ('Ограничения', 'Источники — открытый веб на национальных языках. Аудит полок и платные базы (Mintel GNPD, Euromonitor) не использовались. '
                            'Даты без дня указаны с точностью до месяца или года; «—» — дата не установлена.'),
            ('Оговорка', 'Справочный материал, не юридическая консультация. Перед решениями проверяйте текст нормы по ссылке.'),
        ],
        'heads': ('Статусы', 'Типы мер'),
    },
    'en': {
        'bloc': {'EU': 'EU', 'EAEU': 'EAEU', 'other': 'Outside EU and EAEU'},
        'use': {'high': 'High', 'medium': 'Medium', 'low': 'Low', 'none': 'None', 'nodata': 'No data'},
        'type': {
            'disclosure': 'Ingredient disclosure / labelling', 'claims': 'Claim rules',
            'ban': 'Composition ban / restriction', 'dairy': 'Dairy and milk fat substitutes',
            'import': 'Import control', 'traceability': 'Digital marking and traceability',
            'tax': 'Taxes', 'standard': 'Quality standards', 'sustainability': 'Sustainability (EUDR)',
        },
        'status': {'in_force': 'In force', 'upcoming': 'Adopted, applies later', 'proposal': 'Proposal / draft'},
        'scope': {'EU': 'EU (common rules)', 'EAEU': 'EAEU (common rules)'},
        'sheets': ('Country summary', 'Rules', 'Reference'),
        'mcols': [('No.', 5), ('Scope code', 9), ('Scope', 20), ('Type of rule', 24), ('Status', 17), ('Applies from', 13),
                  ('Legal act / source of the rule', 34), ('Rule', 34), ('Description', 70), ('Link', 40)],
        'ccols': [('No.', 5), ('Country', 16), ('Bloc', 14), ('On-pack wording', 26),
                  ('Use of the “palm oil free” claim', 16), ('Rules total (incl. bloc)', 11), ('In force', 11),
                  ('Apply later', 11), ('Proposals', 11), ('Rules in brief', 52), ('Market practice of the claim', 60),
                  ('Code', 7), ('Bloc code', 8)],
        'total': 'Total rules in the database',
        'note': 'Columns F–I are formulas over the “Rules” sheet: the country’s own rules plus the common rules of its bloc (EU or EAEU). '
                'Column E is an expert estimate based on published sources, chosen from a list.',
        'guide': [
            ('Regulatory policy: labelling and palm oil', ''),
            ('Prepared by', 'International Marketing Agency (IMA)'),
            ('Data as of', None), ('Review period', None), ('', ''), ('Statuses', ''), 'STATUS', ('', ''), ('Types of rules', ''), 'TYPE', ('', ''),
            ('Claim-use scale', 'Expert estimate based on published sources: High / Medium / Low / None / No data.'),
            ('Limitations', 'Sources: the open web in national languages. Shelf audits and paid databases (Mintel GNPD, Euromonitor) were not used. '
                            'Dates without a day are given to the month or year; “—” means no date has been set.'),
            ('Disclaimer', 'Reference material, not legal advice. Check the text of the rule via the link before making decisions.'),
        ],
        'heads': ('Statuses', 'Types of rules'),
    },
}
args = [a for a in sys.argv[1:] if not a.startswith('--lang')]
LANG = 'en' if '--lang=en' in sys.argv[1:] else 'ru'
T = L[LANG]
OUT = Path(args[0]) if args else ROOT / 'public' / 'lab' / ('regulation-en.xlsx' if LANG == 'en' else 'regulation.xlsx')
BLOC, USE, TYPE, STATUS = T['bloc'], T['use'], T['type'], T['status']
SCOPE_NAME = dict(T['scope'])
SCOPE_NAME.update({c['code']: c['name'][LANG] for c in DATA['countries']})

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='17241F')
HEAD_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
LINK = Font(name=FONT, size=10, color='0563C1', underline='single')
THIN = Side(style='thin', color='D9DFD6')
GRID = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
STATUS_FILL = {
    STATUS['in_force']: PatternFill('solid', fgColor='E5F4E4'),
    STATUS['upcoming']: PatternFill('solid', fgColor='F9EEDA'),
    STATUS['proposal']: PatternFill('solid', fgColor='F0EFEC'),
}


def fmt_date(d: str) -> str:
    """ISO-дата → ДД.ММ.ГГГГ; «2025-11» → 11.2025; «2025» → 2025; пусто → «—»."""
    if not d:
        return '—'
    p = d.split('-')
    return '.'.join(reversed(p))


def header(ws, cols):
    for i, (title, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=title)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, Alignment(wrap_text=True, vertical='center'), GRID
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'C2'


def body(cell, font=BODY):
    cell.font, cell.alignment, cell.border = font, WRAP, GRID


wb = Workbook()

# ---------- Лист «Меры»: одна строка — одна норма ----------
ms = wb.active
ms.title = T['sheets'][1]
header(ms, T['mcols'])
NCOL = len(T['mcols'])
for r, m in enumerate(DATA['measures'], 2):
    vals = [r - 1, m['scope'], SCOPE_NAME[m['scope']], TYPE[m['type']], STATUS[m['status']], fmt_date(m['date']),
            m['act'][LANG], m['title'][LANG], m['text'][LANG]]
    if LANG == 'ru':
        vals += [m['title']['en'], m['text']['en']]
    vals.append(m['url'] or '—')
    for col, v in enumerate(vals, 1):
        c = ms.cell(row=r, column=col, value=v)
        body(c, BOLD if col == 8 else BODY)
    ms.cell(row=r, column=5).fill = STATUS_FILL[STATUS[m['status']]]
    if m['url']:
        link = ms.cell(row=r, column=NCOL)
        link.hyperlink, link.font = m['url'], LINK
last_m = len(DATA['measures']) + 1
ms.auto_filter.ref = f'A1:{get_column_letter(NCOL)}{last_m}'

# ---------- Лист «Сводка по странам» ----------
ss = wb.create_sheet(T['sheets'][0], 0)
header(ss, T['ccols'])
M = f"'{T['sheets'][1]}'"
for r, c in enumerate(DATA['countries'], 2):
    bloc_code = c['bloc'] if c['bloc'] in ('EU', 'EAEU') else '-'
    phrase = c['claimPhrase']
    if LANG == 'en' and phrase != '—':
        phrase += ' (“palm oil free”)'
    vals = [r - 1, c['name'][LANG], BLOC[c['bloc']], phrase, USE[c['claimUse']]]
    for col, v in enumerate(vals, 1):
        body(ss.cell(row=r, column=col, value=v), BOLD if col == 2 else BODY)
    # Счётчики норм — формулами по листу «Меры»: свои нормы страны + общие нормы её блока.
    rng_code = f'{M}!$B$2:$B${last_m}'
    rng_status = f'{M}!$E$2:$E${last_m}'
    ss.cell(row=r, column=6, value=f'=COUNTIF({rng_code},L{r})+COUNTIF({rng_code},M{r})')
    for col, st in ((7, STATUS['in_force']), (8, STATUS['upcoming']), (9, STATUS['proposal'])):
        ss.cell(row=r, column=col, value=f'=COUNTIFS({rng_code},L{r},{rng_status},"{st}")+COUNTIFS({rng_code},M{r},{rng_status},"{st}")')
    for col in range(6, 10):
        cell = ss.cell(row=r, column=col)
        body(cell)
        cell.alignment = Alignment(horizontal='center', vertical='top')
    body(ss.cell(row=r, column=10, value=c['summary'][LANG]))
    body(ss.cell(row=r, column=11, value=c['practice'][LANG]))
    body(ss.cell(row=r, column=12, value=c['code']))
    body(ss.cell(row=r, column=13, value=bloc_code))
last_c = len(DATA['countries']) + 1
ss.auto_filter.ref = f'A1:M{last_c}'
dv = DataValidation(type='list', formula1='"' + ','.join(USE.values()) + '"', allow_blank=True)
ss.add_data_validation(dv)
dv.add(f'E2:E{last_c}')
tot = last_c + 2
ss.cell(row=tot, column=2, value=T['total']).font = BOLD
ss.cell(row=tot, column=6, value=f'=COUNTA({M}!$A$2:$A${last_m})').font = BOLD
for col, st in ((7, STATUS['in_force']), (8, STATUS['upcoming']), (9, STATUS['proposal'])):
    ss.cell(row=tot, column=col, value=f'=COUNTIF({M}!$E$2:$E${last_m},"{st}")').font = BOLD
ss.cell(row=tot + 1, column=2,
        value=T['note']).font = Font(name=FONT, size=9, italic=True)

# ---------- Лист «Справочник» ----------
gs = wb.create_sheet(T['sheets'][2])
gs.column_dimensions['A'].width = 30
gs.column_dimensions['B'].width = 90
rows = []
for item in T['guide']:
    if item == 'STATUS':
        rows += [(v, '') for v in STATUS.values()]
    elif item == 'TYPE':
        rows += [(v, '') for v in TYPE.values()]
    elif item[1] is None:
        rows.append((item[0], fmt_date(DATA['updatedAt']) if len(rows) == 2 else DATA['period'][LANG]))
    else:
        rows.append(item)
for r, (a, b) in enumerate(rows, 1):
    ca, cb = gs.cell(row=r, column=1, value=a), gs.cell(row=r, column=2, value=b)
    ca.font = Font(name=FONT, size=13, bold=True) if r == 1 else (BOLD if a in T['heads'] or b else BODY)
    cb.font, cb.alignment = BODY, WRAP
for v, fill in STATUS_FILL.items():
    for r in range(1, len(rows) + 1):
        if gs.cell(row=r, column=1).value == v:
            gs.cell(row=r, column=1).fill = fill

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
